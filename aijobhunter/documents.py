"""Document text extraction for resumes and other inputs.

Reads plain text, Markdown, PDF, and DOCX into a single string. PDF/DOCX support
is what lets ``resume_file`` point at a real CV (``.pdf`` / ``.docx``) instead of
a hand-pasted text blob. Extension drives the reader; unknown types fall back to
UTF-8 text.
"""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


class DocumentReadError(Exception):
    """Raised when a document cannot be read."""


def read_document(path: str | Path) -> str:
    """Return the text content of a document, dispatching on file extension."""
    p = Path(path)
    if not p.exists():
        raise DocumentReadError(f"Document not found: {p}")

    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(p)
    if suffix == ".docx":
        return _read_docx(p)
    # .txt, .md, and anything else: treat as UTF-8 text.
    return p.read_text(encoding="utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise DocumentReadError(
            "pypdf is required to read PDF files (pip install -r requirements.txt)"
        ) from exc
    try:
        reader = PdfReader(str(path))
        parts = [(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:  # noqa: BLE001 - surface a clean error
        raise DocumentReadError(f"Failed to read PDF {path}: {exc}") from exc
    return "\n".join(parts).strip()


def _read_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise DocumentReadError(
            "python-docx is required to read DOCX files (pip install -r requirements.txt)"
        ) from exc
    try:
        document = docx.Document(str(path))
        parts = [para.text for para in document.paragraphs]
        # Include table cell text — resumes often use tables for layout.
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
    except Exception as exc:  # noqa: BLE001
        raise DocumentReadError(f"Failed to read DOCX {path}: {exc}") from exc
    return "\n".join(p for p in parts if p and p.strip()).strip()
